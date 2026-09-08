from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ABSTRACT_RESULTS = (
    "We validated the SDK along five axes. A zero-install reproducibility walk-through took the SDK from browser "
    "page-load to a rendered report in 1.89 s. In-browser adapter execution numerically agreed with conventional "
    "local execution of each wrapped tool in a Lung-AdenoCA SBS96 benchmark, with exact reproduction for three of "
    "the four adapters and floating-point-noise agreement for the fourth. In a separate four-tool comparison, all "
    "adapters analyzed the same spectra and catalog under an explicit harmonization contract. On a published "
    "known-truth benchmark, mean raw exposure cosine ranged from 0.9983 to 1.0000 at 0% noise, 0.9879 to 0.9996 at "
    "5% noise, and 0.9719 to 0.9926 at 10% noise. On 38 real PCAWG Lung-AdenoCA spectra, pairwise exposure "
    "correlations ranged from 0.803 to 0.953 and mean active-signature Jaccard similarity ranged from 0.186 to "
    "0.694, demonstrating that the common adapter layer can identify tool-sensitive conclusions. The SDK's "
    "internal solvers agreed numerically with established reference implementations. Runtime benchmarks across "
    "Chrome, Edge, and Firefox showed that a 120-sample cohort NNLS exposure refit (fitting only, excluding "
    "bootstrap uncertainty) completed in approximately 25 ms, and a 300-sample, 40-signature refit in under 3 s."
)

EXTERNAL_ORCHESTRATION = (
    "The adapter layer lets a single browser session execute and compare multiple established mutational-signature "
    "tools on shared spectra. Each adapter follows a uniform three-stage contract: an input-preparation stage "
    "produces virtual files and a manifest in the target tool's expected matrix orientation; an execution stage runs "
    "the tool in the chosen runtime; and an output-parsing stage reimports the tool's signature and exposure tables "
    "into SDK matrix objects through a shared parser. The common parser preserves the same explicit context order "
    "and catalog labels, but outputs are not assumed to be semantically comparable merely because they share a "
    "matrix shape. Cross-tool comparisons use a separate harmonization contract that records exposure units, "
    "complete-catalog membership, filtering order, inactive and omitted signatures, unassigned components, and "
    "reconstruction normalization. Each adapter records the targeted package version in provenance, alongside an "
    "internal adapter schema version, so downstream consumers can detect schema changes. Four pinned adapter "
    "workflows execute in-browser today (Table 1). These adapters expose selected supervised-fitting or refitting "
    "interfaces from the upstream packages rather than full browser replicas of each package's complete API."
)

VALIDATION_OVERVIEW = (
    "We validated mSigSDK along five axes: that the SDK can be loaded and used without installation; that each "
    "in-browser adapter reproduces its matched local package execution; that a harmonized four-tool comparison "
    "quantifies known-truth accuracy and real-cohort disagreement; that the SDK's internal solvers reproduce "
    "established reference implementations; and that the SDK performs well enough to be used interactively in "
    "common cohort-analysis scenarios."
)

METHODS_PARAGRAPHS = [
    (
        "For cross-tool evaluation, all four adapters received the same ordered SBS96 spectra matrix and the same "
        "signature catalog. After adapter-specific parsing, output columns were mapped to the complete catalog; "
        "package-omitted catalog signatures were retained with zero exposure, non-finite or negative comparison "
        "values were set to zero, and package-native unassigned or unexplained components were not relabeled as "
        "catalog signatures. Each complete output vector was then converted to relative fractions. Raw exposure "
        "accuracy was calculated before filtering. Separately, a common 1% relative-exposure reporting cutoff was "
        "applied, values below the cutoff were set to zero, and retained values were renormalized. Active-signature "
        "counts, precision, recall, F1, and pairwise Jaccard similarity were calculated from these thresholded "
        "vectors. For reconstruction metrics, every catalog signature column was normalized to unit L1 sum before "
        "reconstructing the normalized observed spectrum. Pairwise comparisons used flattened-vector Pearson "
        "correlation and per-sample L1 distance. Metrics were computed from unrounded fractions; rounding was used "
        "only for presentation."
    ),
    (
        "The known-truth analysis used the peer-reviewed synthetic SBS benchmark generated by Islam et al. and "
        "distributed with the Diaz-Gay et al. benchmark archive. The archive contains 2,700 spectra generated from "
        "21 reference signatures across nine cancer-type groups and 0%, 5%, and 10% noise conditions. We analyzed "
        "one archived spectrum from each cancer-type group at each noise level (nine spectra per noise level; 27 "
        "spectra total). The archived 78-signature SBS catalog was supplied unchanged to all four tools. Complete "
        "relative-exposure vectors were compared with the published true fractions by cosine similarity, root-mean-"
        "square error, and mean absolute error; thresholded active calls were compared with the published true "
        "activities by precision, recall, and F1. This high-burden controlled subset was used as an accuracy and "
        "noise-sensitivity test, not as a mutation-burden sweep."
    ),
    (
        "The real-world triangulation analysis used the same 38 PCAWG Lung-AdenoCA WGS SBS96 spectra and the same "
        "67-signature COSMIC v3 GRCh37 SBS96 catalog used for adapter-fidelity testing. All tools used the common "
        "comparison seed 104729 and the package-specific options recorded in the benchmark provenance. The "
        "harmonization layer standardized labels, catalog membership, units, filtering order, and comparison "
        "metrics; it did not make the packages' optimization objectives, sparsity behavior, internal normalization, "
        "signature-selection logic, or native post-processing equivalent. Because the real cohort has no known "
        "exposure truth, this analysis quantifies disagreement and reconstruction behavior rather than biological "
        "accuracy or tool superiority."
    ),
]

RESULTS_PARAGRAPHS = [
    (
        "On the published known-truth benchmark, all four tools closely recovered the true exposure fractions, with "
        "degradation under the archived noise conditions (Figure 5A-B). Mean raw exposure cosine at 0% noise was "
        "0.9983 for deconstructSigs, 1.0000 for sigminer, 1.0000 for SigProfilerAssignment, and 0.9987 for MuSiCal; "
        "at 5% noise it was 0.9879, 0.9919, 0.9996, and 0.9984; and at 10% noise it was 0.9719, 0.9804, 0.9894, and "
        "0.9926, respectively. After the common 1% reporting cutoff, mean active-call F1 at 10% noise was 0.5711, "
        "0.5644, 0.9568, and 0.9167, respectively. At that noise level, mean precision/recall was 0.4459/0.8635 for "
        "deconstructSigs, 0.4474/0.8646 for sigminer, 1.0000/0.9238 for SigProfilerAssignment, and 0.9286/0.9175 for "
        "MuSiCal. Thus, high reconstruction agreement did not eliminate meaningful differences in false-positive "
        "and missed active-signature calls."
    ),
    (
        "The real PCAWG comparison showed materially larger tool-to-tool differences after the same harmonization "
        "steps (Figure 5C-F; Table 2). Flattened exposure correlations ranged from 0.803 for sigminer versus "
        "SigProfilerAssignment to 0.953 for deconstructSigs versus sigminer. Mean per-sample L1 disagreement ranged "
        "from 0.350 to 0.907, and mean active-signature Jaccard similarity ranged from 0.186 to 0.694. Mean active-"
        "signature counts were 15.87 for deconstructSigs, 17.39 for sigminer, 4.13 for SigProfilerAssignment, and "
        "11.21 for MuSiCal. Mean reconstruction cosine remained high for all four tools (0.982-0.995), showing that "
        "similar reconstruction quality can coexist with different exposure allocations and active-signature sets. "
        "The largest mean across-tool exposure ranges were concentrated in SBS5, SBS40, SBS8, SBS4, SBS39, and "
        "SBS18."
    ),
    (
        "The synthetic and real-cohort analyses answer different questions. The synthetic benchmark supports "
        "known-truth accuracy and active-call sensitivity under its archived noise conditions, but its high-burden "
        "mixtures were generated from the same catalog used for fitting and do not demonstrate that the tools are "
        "interchangeable. The PCAWG analysis has no known exposure truth and therefore cannot identify a biologically "
        "superior tool. Instead, it demonstrates the practical value of the common adapter layer: investigators can "
        "distinguish findings that are stable across fitting implementations from findings that are tool-sensitive "
        "and require cautious interpretation."
    ),
]

FIGURE_CAPTION = (
    "Figure 5. Known-truth accuracy and real-world tool triangulation after explicit harmonization. Panels A-B use "
    "one archived spectrum from each of nine cancer-type groups at 0%, 5%, and 10% noise. (A) Mean cosine between "
    "the complete raw relative-exposure vector and the published true fractions, evaluated before filtering. (B) "
    "Mean active-signature F1 after a common 1% relative-exposure cutoff and renormalization. Panels C-F use 38 "
    "PCAWG Lung-AdenoCA WGS SBS96 spectra fitted against the 67-signature COSMIC v3 GRCh37 catalog. (C) Pairwise "
    "Pearson correlation of flattened complete-catalog relative fractions. (D) Mean per-sample L1 distance between "
    "relative-exposure vectors. (E) Mean active-signature Jaccard similarity after the common 1% cutoff. (F) The six "
    "signatures with the largest mean across-tool exposure range. Package-omitted catalog columns were zero-filled; "
    "non-finite or negative comparison values were set to zero; unassigned components were not relabeled as catalog "
    "signatures; and catalog columns were normalized to unit sum for reconstruction metrics. The synthetic panels "
    "evaluate accuracy under the tested archived conditions. The PCAWG panels quantify disagreement without known "
    "truth and do not establish biological superiority."
)

DISCUSSION = (
    "The four-tool benchmark supports a narrower and more useful scientific value proposition than orchestration "
    "alone. On the published synthetic benchmark, all four tools closely recovered the known fractions, although "
    "raw exposure accuracy and active-call performance degraded under increasing archived noise. On the real PCAWG "
    "cohort, explicit harmonization did not eliminate substantial differences in exposure allocation, active-"
    "signature counts, or active-signature concordance, even when reconstruction cosine remained high. The common "
    "adapter layer therefore enables investigators to identify conclusions that are stable across fitting "
    "implementations and conclusions that are tool-sensitive; because the real cohort has no known exposure truth, "
    "the comparison does not establish that one package is biologically superior. The browser-versus-local fidelity "
    "experiment answers a separate question: it shows that the adapters reproduce their wrapped tools rather than "
    "silently altering them. Fidelity does not itself validate the biological correctness of a tool's output. This "
    "numerical result also does not establish security. Browsers are complex, frequently patched attack surfaces, "
    "and JavaScript dependency chains are a documented vector for supply-chain compromise; mSigSDK's pinned-version, "
    "checksum-verified execution and optional no-egress mode reduce but do not eliminate this exposure. Users "
    "handling regulated or sensitive data should weigh these residual risks rather than treat data residency alone "
    "as a privacy control."
)

LIMITATION_OLD = (
    "The present validation establishes browser/local numerical fidelity for a Lung-AdenoCA SBS96 fitting benchmark, "
    "but does not yet validate performance across additional tumor types, low-burden or panel/WES inputs, MAF-"
    "conversion workflows, DBS/ID/SBS1536 profiles, or alternate genome builds."
)

LIMITATION_NEW = (
    "The present validation separates browser/local numerical fidelity for one Lung-AdenoCA SBS96 configuration, "
    "known-truth accuracy for 27 published high-burden synthetic spectra across three archived noise conditions, and "
    "tool disagreement for 38 real Lung-AdenoCA spectra without exposure truth. The synthetic subset is not a "
    "mutation-burden sweep, and the real-cohort comparison cannot determine biological superiority. Performance "
    "across additional tumor types, low-burden or panel/WES inputs, MAF-conversion workflows, DBS/ID/SBS1536 "
    "profiles, and alternate genome builds remains untested."
)

CONCLUSION = (
    "mSigSDK is a browser-native, zero-install, FAIR-aligned JavaScript software development kit for mutational-"
    "signature analysis. It orchestrates four established mutational-signature packages - SigProfilerAssignment, "
    "MuSiCal, deconstructSigs, and sigminer - through a uniform adapter layer; provides exploratory non-negative "
    "matrix factorization in the browser; scores fits against a configurable, burden-aware quality-control evidence "
    "panel; and serializes results under a proposed, versioned JSON Schema with full provenance. In an explicitly "
    "harmonized four-tool benchmark, all four tools showed high known-truth exposure accuracy under the tested "
    "published synthetic conditions, while the real PCAWG cohort showed material differences in exposure allocation "
    "and active-signature calls. These results support multi-tool triangulation as a way to identify stable and tool-"
    "sensitive conclusions, not as a method for selecting a biologically superior tool without truth data. The SDK "
    "loads without installation, runs entirely on the client for the workflows that should be local, completes a full "
    "single-sample analysis against a 67-signature COSMIC catalog in under two seconds from a cold browser load, and "
    "reproduces the behavior of the wrapped tools to numerical precision. It is accessible to anyone working with "
    "mutational signatures - including computational analysts, clinical and translational researchers, methods "
    "developers, and educators - without the configuration overhead of conventional installations."
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def find_paragraph(doc: Document, starts_with: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(starts_with):
            return paragraph
    raise ValueError(f"Paragraph not found: {starts_with}")


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Arial"


def make_summary_table(doc: Document):
    rows = [
        ("deconstructSigs", "15.87", "0.9941", "0.9811"),
        ("sigminer", "17.39", "0.9945", "0.9831"),
        ("SigProfilerAssignment", "4.13", "0.9815", "0.9430"),
        ("MuSiCal", "11.21", "0.9931", "0.9823"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.15), Inches(1.35), Inches(1.45), Inches(1.45)]
    headers = ["Tool", "Mean active signatures", "Mean reconstruction cosine", "Minimum reconstruction cosine"]
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[index]
        cell.text = header
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
    for row_values in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row_values)):
            cell.width = widths[index]
            cell.text = value
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
    return table


def build_revised_manuscript(source: Path, figure_png: Path, output: Path) -> None:
    doc = Document(source)
    replace_paragraph(find_paragraph(doc, "We validated the SDK along four axes."), ABSTRACT_RESULTS)
    replace_paragraph(find_paragraph(doc, "The adapter layer lets a single browser session"), EXTERNAL_ORCHESTRATION)
    replace_paragraph(find_paragraph(doc, "We validated mSigSDK along four axes:"), VALIDATION_OVERVIEW)

    native_heading = find_paragraph(doc, "Native signature fitting")
    heading = native_heading.insert_paragraph_before("Cross-tool comparison and harmonization")
    heading.style = "Heading 3"
    for text in METHODS_PARAGRAPHS:
        paragraph = native_heading.insert_paragraph_before(text)
        paragraph.style = "normal"

    discussion_heading = find_paragraph(doc, "Discussion")
    new_heading = discussion_heading.insert_paragraph_before("Known-truth accuracy and real-world tool triangulation")
    new_heading.style = "Heading 3"
    for text in RESULTS_PARAGRAPHS:
        paragraph = discussion_heading.insert_paragraph_before(text)
        paragraph.style = "normal"
    table_caption = discussion_heading.insert_paragraph_before(
        "Table 2. Active-signature counts and reconstruction quality for the real PCAWG four-tool comparison after the common 1% reporting cutoff."
    )
    table_caption.style = "normal"
    table_caption.runs[0].bold = True
    table = make_summary_table(doc)
    discussion_heading._p.addprevious(table._tbl)
    figure_paragraph = discussion_heading.insert_paragraph_before()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.add_run().add_picture(str(figure_png), width=Inches(6.5))
    caption = discussion_heading.insert_paragraph_before(FIGURE_CAPTION)
    caption.style = "normal"
    caption.runs[0].bold = True

    replace_paragraph(find_paragraph(doc, "The validation results carry three substantive implications."), DISCUSSION)
    limitations = find_paragraph(doc, "Several limitations warrant consideration.")
    if LIMITATION_OLD not in limitations.text:
        raise ValueError("Expected limitations sentence not found")
    replace_paragraph(limitations, limitations.text.replace(LIMITATION_OLD, LIMITATION_NEW))
    replace_paragraph(find_paragraph(doc, "mSigSDK is a browser-native, zero-install, FAIR-aligned JavaScript"), CONCLUSION)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def style_guide(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in (
        ("Title", 22, "172B3A", 0, 8),
        ("Heading 1", 16, "2E74B5", 14, 7),
        ("Heading 2", 13, "2E74B5", 11, 5),
        ("Heading 3", 11, "1F4D78", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_labeled_block(doc: Document, label: str, text: str, fill: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    label_run = paragraph.add_run(f"{label}\n")
    label_run.bold = True
    label_run.font.name = "Arial"
    label_run.font.size = Pt(9)
    body = paragraph.add_run(text)
    body.font.name = "Arial"
    body.font.size = Pt(9)


def build_guide(source: Path, figure_png: Path, output: Path) -> None:
    original = Document(source)
    originals = {
        "abstract": find_paragraph(original, "We validated the SDK along four axes.").text,
        "orchestration": find_paragraph(original, "The adapter layer lets a single browser session").text,
        "overview": find_paragraph(original, "We validated mSigSDK along four axes:").text,
        "discussion": find_paragraph(original, "The validation results carry three substantive implications.").text,
        "limitations": LIMITATION_OLD,
        "conclusion": find_paragraph(original, "mSigSDK is a browser-native, zero-install, FAIR-aligned JavaScript").text,
    }
    doc = Document()
    style_guide(doc)
    title = doc.add_paragraph("Reviewer 4 - Major Comment 1 manuscript edit kit", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph("Exact changes for the central four-tool value proposition")
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("526273")

    doc.add_heading("What to do", level=1)
    add_labeled_block(
        doc,
        "FIGURE DECISION",
        "Do not replace Figures 1-4 and do not reuse the adapter-fidelity figure. Add the new benchmark as Figure 5 after current Figure 4, immediately before the Discussion. The adapter-fidelity experiment remains in place because it answers browser-versus-local fidelity, not cross-tool scientific triangulation.",
        "EAF4F8",
    )
    for item in (
        "Replace three existing overview paragraphs: Abstract Results, External tool orchestration, and the opening Results overview.",
        "Insert one Methods subsection after Table 1 and before Native signature fitting.",
        "Insert one Results subsection, Table 2, Figure 5, and its caption after current Figure 4.",
        "Replace the first Discussion implication and narrowly update the limitations and Conclusions.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    edits = [
        ("1. Abstract - Results paragraph", originals["abstract"], ABSTRACT_RESULTS),
        ("2. Implementation - External tool orchestration", originals["orchestration"], EXTERNAL_ORCHESTRATION),
        ("3. Results - validation overview", originals["overview"], VALIDATION_OVERVIEW),
    ]
    doc.add_heading("Exact replacements", level=1)
    for heading, old, new in edits:
        doc.add_heading(heading, level=2)
        add_labeled_block(doc, "ORIGINAL", old, "F2F4F7")
        add_labeled_block(doc, "REPLACE WITH", new, "EAF4F8")

    doc.add_heading("Exact Methods insertion", level=1)
    doc.add_paragraph("Placement: after Table 1 and before the heading Native signature fitting.")
    doc.add_heading("Cross-tool comparison and harmonization", level=2)
    for paragraph in METHODS_PARAGRAPHS:
        doc.add_paragraph(paragraph)
    add_labeled_block(
        doc,
        "WHY THIS SATISFIES THE REVIEWER",
        "This insertion explicitly defines absolute-to-relative conversion, complete-catalog zero filling, inactive and unassigned components, the 1% cutoff and renormalization order, catalog normalization, numerical handling, active-call metrics, reconstruction quality, pairwise exposure correlation, per-sample L1 disagreement, and the boundary between harmonized reporting semantics and package-specific algorithms.",
        "FFF4D6",
    )

    doc.add_heading("Exact Results insertion", level=1)
    doc.add_paragraph("Placement: after current Figure 4 and before Discussion. This preserves existing figure numbering and makes the new benchmark Figure 5.")
    doc.add_heading("Known-truth accuracy and real-world tool triangulation", level=2)
    for paragraph in RESULTS_PARAGRAPHS:
        doc.add_paragraph(paragraph)
    caption = doc.add_paragraph("Table 2. Active-signature counts and reconstruction quality for the real PCAWG four-tool comparison after the common 1% reporting cutoff.")
    caption.runs[0].bold = True
    make_summary_table(doc)
    doc.add_paragraph()
    figure_title = doc.add_paragraph("New Figure 5 - add; do not replace an existing figure")
    figure_title.runs[0].bold = True
    figure_title.paragraph_format.keep_with_next = True
    figure_para = doc.add_paragraph()
    figure_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_para.paragraph_format.keep_with_next = True
    figure_para.add_run().add_picture(str(figure_png), width=Inches(6.8))
    add_labeled_block(doc, "FIGURE 5 CAPTION", FIGURE_CAPTION, "EAF4F8")

    doc.add_heading("Discussion, limitations, and Conclusions", level=1)
    for heading, old, new in (
        ("Discussion - first implication", originals["discussion"], DISCUSSION),
        ("Limitations - replace only the quoted sentence", originals["limitations"], LIMITATION_NEW),
        ("Conclusions", originals["conclusion"], CONCLUSION),
    ):
        doc.add_heading(heading, level=2)
        add_labeled_block(doc, "ORIGINAL", old, "F2F4F7")
        add_labeled_block(doc, "REPLACE WITH", new, "EAF4F8")

    doc.add_heading("Final consistency check", level=1)
    checks = [
        "The manuscript now separates adapter fidelity, known-truth accuracy, and real-world triangulation.",
        "Raw accuracy is evaluated before the reporting cutoff; active calls are evaluated after the common 1% cutoff and renormalization.",
        "The synthetic benchmark is not described as a mutation-burden sweep or evidence that tools are interchangeable.",
        "The real PCAWG comparison is not described as accuracy or evidence that one tool is biologically superior.",
        "Figure 5 contains all four real-world displays requested by the reviewer: exposure correlation, per-sample disagreement, active-signature concordance, and discrepancy-driving signatures.",
        "Table 2 reports active-signature counts and reconstruction quality explicitly.",
    ]
    for check in checks:
        doc.add_paragraph(check, style="List Bullet")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--figure-png", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    final_figure = args.output_dir / "figure5-four-tool-known-truth-and-triangulation.png"
    shutil.copyfile(args.figure_png, final_figure)
    build_revised_manuscript(
        args.source,
        final_figure,
        args.output_dir / "mSigSDK-reviewer4-comment1-revised-manuscript.docx",
    )
    build_guide(
        args.source,
        final_figure,
        args.output_dir / "reviewer4-comment1-manuscript-edit-guide.docx",
    )
    print(args.output_dir)


if __name__ == "__main__":
    main()
